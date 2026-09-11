import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference
from pathlib import Path

# ==========================================
# FILE SETTINGS
# ==========================================

folder = Path("C:\\Users\\admin\\Desktop\\decode projects\\project2\\Dataset for Data Analytics (4)").parent

dataset_file = folder / "Dataset for Data Analytics (4).xlsx"
output_file = folder / "DecodeLabs_Project2_Final.xlsx"

monthly_chart = folder / "monthly_revenue_trend.png"
product_chart = folder / "revenue_by_product.png"
distribution_chart = folder / "total_price_distribution.png"
correlation_chart = folder / "correlation_heatmap.png"


# ==========================================
# LOAD DATA
# ==========================================

df = pd.read_excel(dataset_file)

df["Date"] = pd.to_datetime(df["Date"])

numeric_columns = [
    "Quantity",
    "UnitPrice",
    "ItemsInCart",
    "TotalPrice"
]


# ==========================================
# CALCULATE STATISTICS
# ==========================================

statistics = pd.DataFrame({
    "Count": df[numeric_columns].count(),
    "Mean": df[numeric_columns].mean(),
    "Median": df[numeric_columns].median(),
    "Minimum": df[numeric_columns].min(),
    "Maximum": df[numeric_columns].max(),
    "Standard Deviation": df[numeric_columns].std()
})


# ==========================================
# OUTLIER ANALYSIS
# ==========================================

outlier_results = []

for column in numeric_columns:

    Q1 = df[column].quantile(0.25)
    Q3 = df[column].quantile(0.75)
    IQR = Q3 - Q1

    lower_bound = Q1 - 1.5 * IQR
    upper_bound = Q3 + 1.5 * IQR

    outliers = df[
        (df[column] < lower_bound) |
        (df[column] > upper_bound)
    ]

    outlier_results.append([
        column,
        Q1,
        Q3,
        IQR,
        lower_bound,
        upper_bound,
        len(outliers)
    ])

outlier_df = pd.DataFrame(
    outlier_results,
    columns=[
        "Variable",
        "Q1",
        "Q3",
        "IQR",
        "Lower Bound",
        "Upper Bound",
        "Number of Outliers"
    ]
)


# ==========================================
# PRODUCT PERFORMANCE
# ==========================================

product_summary = df.groupby("Product").agg(
    Total_Orders=("OrderID", "count"),
    Total_Quantity=("Quantity", "sum"),
    Total_Revenue=("TotalPrice", "sum"),
    Average_Order_Value=("TotalPrice", "mean")
).sort_values("Total_Revenue", ascending=False)


# ==========================================
# YEARLY SALES
# ==========================================

df["Year"] = df["Date"].dt.year

yearly_sales = df.groupby("Year").agg(
    Total_Orders=("OrderID", "count"),
    Total_Revenue=("TotalPrice", "sum"),
    Average_Order_Value=("TotalPrice", "mean")
)


# ==========================================
# MONTHLY SALES
# ==========================================

df["Month"] = df["Date"].dt.to_period("M")

monthly_sales = df.groupby("Month").agg(
    Total_Orders=("OrderID", "count"),
    Total_Revenue=("TotalPrice", "sum"),
    Average_Order_Value=("TotalPrice", "mean")
)

monthly_sales.index = monthly_sales.index.astype(str)


# ==========================================
# CORRELATION
# ==========================================

correlation = df[numeric_columns].corr()


# ==========================================
# CREATE WORKBOOK
# ==========================================

wb = Workbook()

# Remove default sheet
default_sheet = wb.active
wb.remove(default_sheet)


# ==========================================
# COLORS / FORMATTING
# ==========================================

header_fill = PatternFill(
    fill_type="solid",
    fgColor="1F4E78"
)

section_fill = PatternFill(
    fill_type="solid",
    fgColor="D9EAF7"
)

white_font = Font(
    color="FFFFFF",
    bold=True
)

title_font = Font(
    size=20,
    bold=True
)

subtitle_font = Font(
    size=12,
    italic=True
)

bold_font = Font(
    bold=True
)

thin_border = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin")
)


# ==========================================
# SHEET 1 — RAW DATA
# ==========================================

raw_ws = wb.create_sheet("Raw Data")

for row in df.drop(columns=["Year", "Month"]).itertuples(index=False):
    raw_ws.append(row)

# Add column headers
for col_num, column_name in enumerate(
    df.drop(columns=["Year", "Month"]).columns,
    start=1
):
    cell = raw_ws.cell(row=1, column=col_num)
    cell.value = column_name
    cell.fill = header_fill
    cell.font = white_font
    cell.alignment = Alignment(horizontal="center")

# Format dates
date_column = list(
    df.drop(columns=["Year", "Month"]).columns
).index("Date") + 1

for row in range(2, raw_ws.max_row + 1):
    raw_ws.cell(row=row, column=date_column).number_format = "yyyy-mm-dd"

raw_ws.freeze_panes = "A2"
raw_ws.auto_filter.ref = raw_ws.dimensions


# ==========================================
# SHEET 2 — EDA SUMMARY
# ==========================================

ws = wb.create_sheet("EDA Summary")

ws["A1"] = "DecodeLabs — Project 2"
ws["A1"].font = title_font

ws["A2"] = "Exploratory Data Analysis Report"
ws["A2"].font = subtitle_font

ws["A4"] = "Dataset Overview"
ws["A4"].font = Font(size=14, bold=True)
ws["A4"].fill = section_fill

overview = [
    ["Metric", "Value"],
    ["Total Records", len(df)],
    ["Total Columns", len(df.columns) - 2],
    ["Missing CouponCode Values", df["CouponCode"].isna().sum()],
    ["Duplicate Rows", df.duplicated().sum()],
    ["Earliest Date", df["Date"].min().strftime("%Y-%m-%d")],
    ["Latest Date", df["Date"].max().strftime("%Y-%m-%d")]
]

for row in overview:
    ws.append(row)

for cell in ws[5]:
    cell.fill = header_fill
    cell.font = white_font


# ==========================================
# DESCRIPTIVE STATISTICS
# ==========================================

start_row = 14

ws.cell(start_row, 1).value = "Descriptive Statistics"
ws.cell(start_row, 1).font = Font(size=14, bold=True)
ws.cell(start_row, 1).fill = section_fill

headers = [
    "Variable",
    "Count",
    "Mean",
    "Median",
    "Minimum",
    "Maximum",
    "Standard Deviation"
]

for col, header in enumerate(headers, start=1):
    cell = ws.cell(start_row + 1, col)
    cell.value = header
    cell.fill = header_fill
    cell.font = white_font

for row_num, (variable, row_data) in enumerate(
    statistics.iterrows(),
    start=start_row + 2
):

    ws.cell(row_num, 1).value = variable

    for col_num, value in enumerate(row_data, start=2):
        ws.cell(row_num, col_num).value = value


# ==========================================
# PRODUCT PERFORMANCE
# ==========================================

product_start = 22

ws.cell(product_start, 1).value = "Product Performance"
ws.cell(product_start, 1).font = Font(size=14, bold=True)
ws.cell(product_start, 1).fill = section_fill

product_headers = [
    "Product",
    "Total Orders",
    "Total Quantity",
    "Total Revenue",
    "Average Order Value"
]

for col, header in enumerate(product_headers, start=1):
    cell = ws.cell(product_start + 1, col)
    cell.value = header
    cell.fill = header_fill
    cell.font = white_font

for row_num, (product, row_data) in enumerate(
    product_summary.iterrows(),
    start=product_start + 2
):

    ws.cell(row_num, 1).value = product

    for col_num, value in enumerate(row_data, start=2):
        ws.cell(row_num, col_num).value = value


# ==========================================
# YEARLY SALES
# ==========================================

year_start = 33

ws.cell(year_start, 1).value = "Yearly Sales Trend"
ws.cell(year_start, 1).font = Font(size=14, bold=True)
ws.cell(year_start, 1).fill = section_fill

year_headers = [
    "Year",
    "Total Orders",
    "Total Revenue",
    "Average Order Value"
]

for col, header in enumerate(year_headers, start=1):
    cell = ws.cell(year_start + 1, col)
    cell.value = header
    cell.fill = header_fill
    cell.font = white_font

for row_num, (year, row_data) in enumerate(
    yearly_sales.iterrows(),
    start=year_start + 2
):

    ws.cell(row_num, 1).value = year

    for col_num, value in enumerate(row_data, start=2):
        ws.cell(row_num, col_num).value = value


# ==========================================
# OUTLIER ANALYSIS
# ==========================================

outlier_start = 40

ws.cell(outlier_start, 1).value = "Outlier Analysis"
ws.cell(outlier_start, 1).font = Font(size=14, bold=True)
ws.cell(outlier_start, 1).fill = section_fill

for col, header in enumerate(outlier_df.columns, start=1):
    cell = ws.cell(outlier_start + 1, col)
    cell.value = header
    cell.fill = header_fill
    cell.font = white_font

for row_num, row_data in enumerate(
    outlier_df.itertuples(index=False),
    start=outlier_start + 2
):

    for col_num, value in enumerate(row_data, start=1):
        ws.cell(row_num, col_num).value = value


# ==========================================
# FORMAT EDA SUMMARY
# ==========================================

for row in ws.iter_rows():
    for cell in row:
        cell.alignment = Alignment(
            vertical="center"
        )

for row in ws.iter_rows():
    for cell in row:
        if cell.value is not None:
            cell.border = thin_border

for column in range(1, 9):
    ws.column_dimensions[get_column_letter(column)].width = 22

ws.freeze_panes = "A5"


# ==========================================
# SHEET 3 — VISUALIZATIONS
# ==========================================

viz = wb.create_sheet("Visualizations")

viz["A1"] = "DecodeLabs — Project 2 Visualizations"
viz["A1"].font = title_font

viz["A3"] = "Monthly Revenue Trend"
viz["A3"].font = Font(size=14, bold=True)

if monthly_chart.exists():
    img = Image(str(monthly_chart))
    img.width = 800
    img.height = 400
    viz.add_image(img, "A4")


viz["A27"] = "Revenue by Product"
viz["A27"].font = Font(size=14, bold=True)

if product_chart.exists():
    img = Image(str(product_chart))
    img.width = 700
    img.height = 420
    viz.add_image(img, "A28")


viz["A55"] = "Total Price Distribution"
viz["A55"].font = Font(size=14, bold=True)

if distribution_chart.exists():
    img = Image(str(distribution_chart))
    img.width = 700
    img.height = 420
    viz.add_image(img, "A56")


viz["A83"] = "Correlation Heatmap"
viz["A83"].font = Font(size=14, bold=True)

if correlation_chart.exists():
    img = Image(str(correlation_chart))
    img.width = 650
    img.height = 500
    viz.add_image(img, "A84")


# ==========================================
# SHEET 4 — KEY FINDINGS
# ==========================================

findings = wb.create_sheet("Key Findings")

findings["A1"] = "DecodeLabs — Project 2"
findings["A1"].font = title_font

findings["A2"] = "Key EDA Findings and Observations"
findings["A2"].font = subtitle_font


key_findings = [
    (
        "1. Dataset Overview",
        "The dataset contains 1,200 records across 14 original columns. "
        "There are no duplicate rows and the numerical variables contain "
        "1,200 valid observations each."
    ),

    (
        "2. Total Price Distribution",
        "TotalPrice has a mean of 1,053.97 and a median of 823.62. "
        "The higher mean compared with the median suggests that some "
        "higher-value orders pull the average upward."
    ),

    (
        "3. Outliers",
        "The IQR method identified 8 outliers in TotalPrice. "
        "These transactions have the maximum quantity of 5 and relatively "
        "high unit prices. They appear to represent legitimate high-value "
        "orders rather than obvious data-entry errors."
    ),

    (
        "4. Product Performance",
        "Chair generated the highest total revenue at 195,620.11, "
        "closely followed by Printer at 195,612.61. Laptop recorded "
        "the highest average order value at approximately 1,110.56."
    ),

    (
        "5. Lowest Product Revenue",
        "Phone generated the lowest total revenue at 151,722.39 and "
        "also had the fewest orders, with 156 transactions."
    ),

    (
        "6. Monthly Revenue Trend",
        "June 2024 recorded the highest monthly revenue at 68,068.54, "
        "while April 2023 recorded the lowest at 27,751.71. "
        "Monthly revenue shows noticeable fluctuations rather than a "
        "consistent upward or downward pattern."
    ),

    (
        "7. Yearly Trend",
        "Total revenue and order volume were highest in 2023 and lower "
        "in 2024. However, 2025 contains only January through June, "
        "so its yearly total should not be directly compared with the "
        "two complete years."
    ),

    (
        "8. Correlation",
        "UnitPrice has the strongest positive correlation with TotalPrice "
        "(r = 0.72), followed by Quantity (r = 0.62). ItemsInCart has "
        "a weaker positive relationship with TotalPrice (r = 0.39)."
    ),

    (
        "9. Analytical Conclusion",
        "The analysis indicates that order value is most strongly associated "
        "with unit price and quantity. Product revenue is relatively close "
        "among the leading products, while monthly performance fluctuates "
        "considerably across the observation period."
    )
]


row = 5

for title, observation in key_findings:

    findings.cell(row, 1).value = title
    findings.cell(row, 1).font = bold_font

    findings.cell(row + 1, 1).value = observation
    findings.cell(row + 1, 1).alignment = Alignment(
        wrap_text=True,
        vertical="top"
    )

    findings.merge_cells(
        start_row=row + 1,
        start_column=1,
        end_row=row + 2,
        end_column=6
    )

    row += 4


findings.column_dimensions["A"].width = 25

for column in range(2, 7):
    findings.column_dimensions[
        get_column_letter(column)
    ].width = 18


# ==========================================
# SAVE FINAL WORKBOOK
# ==========================================

wb.save(output_file)

print("\n==========================================")
print("FINAL EDA REPORT CREATED SUCCESSFULLY")
print("==========================================")
print(f"File: {output_file}")
print("Sheets:")
print("1. Raw Data")
print("2. EDA Summary")
print("3. Visualizations")
print("4. Key Findings")
print("==========================================")